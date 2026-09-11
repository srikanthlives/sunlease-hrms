"""Bulk attendance upload via .xlsx - mirrors bulk_import_service.py's
exact shape (COLUMNS/SAMPLE_ROW, template+reference-sheet download,
per-row db.begin_nested() savepoints). One long-format row per
(employee, date) handles both a single day (template filtered to one
date) and a full month (template pre-populated with one row per employee
per day of the month)."""
import calendar
import io
from datetime import date, datetime

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session

from app.models.models import AttendanceRecord, EmploymentEpisode, ShiftMaster, User
from app.services import attendance_service, employee_service

# (field_key, column_header). Order here is the order columns appear in
# the downloadable template. "*" in the header marks a required field.
COLUMNS = [
    ("employee_number", "Employee Number*"),
    ("date", "Date* (YYYY-MM-DD)"),
    ("check_in", "Check In (HH:MM)"),
    ("check_out", "Check Out (HH:MM)"),
    ("status", "Status (PRESENT/ABSENT/HALF_DAY/ON_LEAVE/WEEKLY_OFF/HOLIDAY)"),
    ("shift_code", "Shift Code"),
    ("remarks", "Remarks"),
]

SAMPLE_ROW = {
    "employee_number": "EMP00200", "date": "2026-01-01", "check_in": "09:00", "check_out": "18:00",
    "status": "PRESENT", "shift_code": "", "remarks": "",
}

STATUS_VALUES = ["PRESENT", "ABSENT", "HALF_DAY", "ON_LEAVE", "WEEKLY_OFF", "HOLIDAY"]


def build_attendance_template_workbook(db: Session, cost_center_id: int | None = None, year: int | None = None, month: int | None = None) -> Workbook:
    wb = Workbook()
    sheet = wb.active
    sheet.title = "Attendance"

    headers = [header for _, header in COLUMNS]
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = Font(bold=True)

    if cost_center_id is not None and year is not None and month is not None:
        _, last_day = calendar.monthrange(year, month)
        month_start, month_end = date(year, month, 1), date(year, month, last_day)
        episodes = employee_service.episodes_in_cost_center_during(db, cost_center_id, month_start, month_end)
        episodes.sort(key=lambda e: e.employee_number)
        for episode in episodes:
            d = month_start
            while d <= month_end:
                sheet.append([episode.employee_number, d.isoformat(), "", "", "", "", ""])
                d = date.fromordinal(d.toordinal() + 1)
    else:
        sheet.append([SAMPLE_ROW.get(key, "") for key, _ in COLUMNS])

    for col_idx, _ in enumerate(COLUMNS, start=1):
        sheet.column_dimensions[sheet.cell(row=1, column=col_idx).column_letter].width = 26

    ref = wb.create_sheet("Reference Values")

    def _dump(row_label, col_letter, names):
        ref[f"{col_letter}1"] = row_label
        ref[f"{col_letter}1"].font = Font(bold=True)
        for i, name in enumerate(names, start=2):
            ref[f"{col_letter}{i}"] = name
        ref.column_dimensions[col_letter].width = 26

    _dump("Shift Codes", "A", [s.code for s in db.query(ShiftMaster).filter(ShiftMaster.is_active.is_(True)).all()])
    _dump("Status Values", "B", STATUS_VALUES)

    return wb


def _cell_str(value) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def _cell_date(value) -> date | None:
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    try:
        return datetime.strptime(str(value).strip(), "%Y-%m-%d").date()
    except ValueError:
        return None


def _cell_time_combine(date_: date, value) -> datetime | None:
    if value is None or value == "" or date_ is None:
        return None
    if isinstance(value, datetime):
        return datetime.combine(date_, value.time())
    text = str(value).strip()
    if not text:
        return None
    for fmt in ("%H:%M", "%H:%M:%S"):
        try:
            parsed = datetime.strptime(text, fmt)
            return datetime.combine(date_, parsed.time())
        except ValueError:
            continue
    return None


def import_attendance_workbook(db: Session, file_bytes: bytes, actor: User) -> dict:
    wb = load_workbook(io.BytesIO(file_bytes), data_only=True)
    sheet = wb["Attendance"] if "Attendance" in wb.sheetnames else wb.active

    header_row = [(_cell_str(c.value) or "") for c in sheet[1]]
    key_by_column_index = {}
    for idx, header in enumerate(header_row):
        for key, expected_header in COLUMNS:
            if header.strip().lower() == expected_header.strip().lower():
                key_by_column_index[idx] = key
                break

    created = 0
    updated = 0
    errors = []

    for row_number, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
        if row is None or all(v is None or str(v).strip() == "" for v in row):
            continue  # skip blank rows

        data = {}
        for idx, value in enumerate(row):
            key = key_by_column_index.get(idx)
            if key:
                data[key] = value

        employee_number = _cell_str(data.get("employee_number"))
        date_ = _cell_date(data.get("date"))

        if not employee_number or not date_:
            errors.append({"row": row_number, "message": "Employee Number and Date are required"})
            continue

        savepoint = db.begin_nested()
        try:
            episode = db.query(EmploymentEpisode).filter(EmploymentEpisode.employee_number == employee_number).first()
            if not episode:
                raise ValueError(f"Employee Number '{employee_number}' not found")

            status_ = _cell_str(data.get("status")) or "PRESENT"
            if status_.upper() not in STATUS_VALUES:
                raise ValueError(f"Invalid status '{status_}' (must be one of {', '.join(STATUS_VALUES)})")
            status_ = status_.upper()

            shift_code = _cell_str(data.get("shift_code"))
            shift_id = None
            if shift_code:
                shift = db.query(ShiftMaster).filter(ShiftMaster.code == shift_code).first()
                if not shift:
                    raise ValueError(f"Shift Code '{shift_code}' not found")
                shift_id = shift.id

            check_in = _cell_time_combine(date_, data.get("check_in"))
            check_out = _cell_time_combine(date_, data.get("check_out"))

            existing = (
                db.query(AttendanceRecord)
                .filter(AttendanceRecord.episode_id == episode.id, AttendanceRecord.date == date_)
                .first()
            )
            was_existing = existing is not None

            attendance_service.mark_attendance(db, episode.id, date_, check_in, check_out, shift_id, status_, actor)

            savepoint.commit()
            if was_existing:
                updated += 1
            else:
                created += 1
        except (ValueError, IntegrityError) as exc:
            savepoint.rollback()
            errors.append({"row": row_number, "message": str(exc.__cause__ or exc) if isinstance(exc, IntegrityError) else str(exc)})

    return {"created": created, "updated": updated, "errors": errors}
