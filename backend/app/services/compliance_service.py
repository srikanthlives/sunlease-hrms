"""Module 4: Statutory Compliance Management - aggregates already-computed
Module 3 PayslipLine amounts (PF/ESI/PT/LWF) plus FullFinalSettlement
gratuity payouts into per-scheme/cost-center/month ComplianceRecord rows,
and produces a best-effort downloadable contribution statement per
scheme. See CLAUDE.md and the Module 3/4 plan for the overall design -
one calculation (Module 3's payroll processing), two consumers (payslips
themselves, and this module's aggregation)."""
import calendar
import datetime as dt

from openpyxl import Workbook
from openpyxl.styles import Font
from sqlalchemy.orm import Session

from app.models.enums import AuditAction
from app.models.models import (
    EmploymentEpisode, FullFinalSettlement, Payslip, PayslipLine, PayrollRun,
    StatutoryInfo, User, ComplianceRecord,
)
from app.services import approval_service, audit_service

# Employee-side component codes per scheme, matching the exact strings
# payroll_service.py writes into PayslipLine.component_code (see
# compute_statutory_deductions). PT and LWF do have an employee-side
# amount too; LWF additionally has an employer-side amount
# (EMPLOYER_LWF) - PT has no employer-side contribution in this repo's
# model, so its EMPLOYER_SCHEME_COMPONENT_CODES entry is intentionally
# absent (employer contribution stays 0 for PT).
SCHEME_COMPONENT_CODES = {
    "PF": "PF",
    "ESI": "ESI",
    "PT": "PT",
    "LWF": "LWF",
}
EMPLOYER_SCHEME_COMPONENT_CODES = {
    "PF": "EMPLOYER_PF",
    "ESI": "EMPLOYER_ESI",
    "LWF": "EMPLOYER_LWF",
}

VALID_SCHEMES = ("PF", "ESI", "PT", "LWF", "GRATUITY")


def _get_or_create_record(db: Session, scheme: str, cost_center_id: int | None, year: int, month: int) -> ComplianceRecord:
    record = (
        db.query(ComplianceRecord)
        .filter(
            ComplianceRecord.scheme == scheme, ComplianceRecord.cost_center_id == cost_center_id,
            ComplianceRecord.year == year, ComplianceRecord.month == month,
        )
        .first()
    )
    if not record:
        record = ComplianceRecord(scheme=scheme, cost_center_id=cost_center_id, year=year, month=month, status="PENDING")
        db.add(record)
        db.flush()
    return record


def _aggregate_payslip_scheme(db: Session, scheme: str, cost_center_id: int | None, year: int, month: int) -> tuple[float, float, int]:
    """Sums PayslipLine amounts for `scheme` across every Payslip belonging
    to PayrollRun(s) for this (year, month) - across all cost-center runs
    if cost_center_id is None, else just that run. Returns
    (total_employee, total_employer, employees_covered)."""
    run_query = db.query(PayrollRun).filter(PayrollRun.year == year, PayrollRun.month == month)
    if cost_center_id is not None:
        run_query = run_query.filter(PayrollRun.cost_center_id == cost_center_id)
    run_ids = [r.id for r in run_query.all()]
    if not run_ids:
        return 0.0, 0.0, 0

    payslip_ids_query = db.query(Payslip.id).filter(Payslip.run_id.in_(run_ids))
    payslip_ids = [p[0] for p in payslip_ids_query.all()]
    if not payslip_ids:
        return 0.0, 0.0, 0

    employee_code = SCHEME_COMPONENT_CODES[scheme]
    employer_code = EMPLOYER_SCHEME_COMPONENT_CODES.get(scheme)

    employee_lines = (
        db.query(PayslipLine)
        .filter(PayslipLine.payslip_id.in_(payslip_ids), PayslipLine.component_code == employee_code, PayslipLine.amount != 0)
        .all()
    )
    total_employee = round(sum(l.amount for l in employee_lines), 2)
    covered_payslip_ids = {l.payslip_id for l in employee_lines}

    total_employer = 0.0
    if employer_code:
        employer_lines = (
            db.query(PayslipLine)
            .filter(PayslipLine.payslip_id.in_(payslip_ids), PayslipLine.component_code == employer_code, PayslipLine.amount != 0)
            .all()
        )
        total_employer = round(sum(l.amount for l in employer_lines), 2)
        covered_payslip_ids |= {l.payslip_id for l in employer_lines}

    employees_covered = (
        db.query(Payslip.episode_id).filter(Payslip.id.in_(covered_payslip_ids)).distinct().count()
        if covered_payslip_ids else 0
    )
    return total_employee, total_employer, employees_covered


def _aggregate_gratuity(db: Session, cost_center_id: int | None, year: int, month: int) -> tuple[float, float, int]:
    """Gratuity has no monthly PayslipLine (it's computed at Full & Final
    Settlement, not monthly payroll - per the plan's design decision), so
    this sums FullFinalSettlement.gratuity_amount for settlements
    finalized (processed_at) within the given (year, month), optionally
    filtered to episodes currently in the given Cost Center. Employee-side
    contribution is always 0 - gratuity is entirely employer-funded."""
    _, last_day = calendar.monthrange(year, month)
    month_start = dt.datetime(year, month, 1)
    month_end = dt.datetime(year, month, last_day, 23, 59, 59)

    settlements = (
        db.query(FullFinalSettlement)
        .filter(
            FullFinalSettlement.status == "COMPLETED",
            FullFinalSettlement.processed_at >= month_start,
            FullFinalSettlement.processed_at <= month_end,
            FullFinalSettlement.gratuity_amount > 0,
        )
        .all()
    )
    if cost_center_id is not None:
        settlements = [
            s for s in settlements
            if approval_service.current_cost_center_id(db, s.episode_id) == cost_center_id
        ]

    total_employer = round(sum(s.gratuity_amount for s in settlements), 2)
    return 0.0, total_employer, len(settlements)


def aggregate_scheme(db: Session, scheme: str, cost_center_id: int | None, year: int, month: int) -> ComplianceRecord:
    """Upserts a ComplianceRecord for (scheme, cost_center_id, year, month).
    Re-aggregating an already-FILED/PAID record refreshes the contribution
    totals/employees_covered but never resets `status`,
    `challan_reference_number` or `filed_date` back to unfiled - a filing
    record must not be silently wiped out by a later recompute."""
    if scheme not in VALID_SCHEMES:
        raise ValueError(f"Unknown scheme: {scheme}. Must be one of {VALID_SCHEMES}")

    if scheme == "GRATUITY":
        total_employee, total_employer, employees_covered = _aggregate_gratuity(db, cost_center_id, year, month)
    else:
        total_employee, total_employer, employees_covered = _aggregate_payslip_scheme(db, scheme, cost_center_id, year, month)

    record = _get_or_create_record(db, scheme, cost_center_id, year, month)
    record.total_employee_contribution = total_employee
    record.total_employer_contribution = total_employer
    record.employees_covered = employees_covered
    # status/challan_reference_number/filed_date deliberately left untouched
    # here - only mark_filed sets those.
    db.add(record)
    db.flush()
    return record


def mark_filed(db: Session, record_id: int, challan_reference_number: str, filed_date: dt.date, remarks: str | None, user: User) -> ComplianceRecord:
    record = db.query(ComplianceRecord).filter(ComplianceRecord.id == record_id).first()
    if not record:
        raise ValueError("Compliance record not found")
    record.status = "FILED"
    record.challan_reference_number = challan_reference_number
    record.filed_date = filed_date
    record.remarks = remarks
    db.add(record)
    audit_service.record(db, "COMPLIANCE_RECORD", record.id, AuditAction.STATUS_CHANGE, user,
                          new_value=f"FILED ref={challan_reference_number} date={filed_date}")
    return record


# ---------------------------------------------------------------------------
# Downloadable statement (best-effort, not government-portal-certified)
# ---------------------------------------------------------------------------

def build_scheme_statement_workbook(db: Session, record: ComplianceRecord) -> Workbook:
    """Builds a best-effort per-employee contribution statement for one
    ComplianceRecord.

    CAVEAT (mirrors the plan's explicit honesty requirement, and salary-app's
    own scope note): this is a best-effort contribution statement only -
    verify the column layout against the current EPFO/ESIC/state PT/LWF
    portal specification before uploading anywhere. It is NOT a
    certified government template, and should not be treated as one.
    """
    wb = Workbook()
    sheet = wb.active
    sheet.title = f"{record.scheme} {record.year}-{record.month:02d}"

    if record.scheme in ("PF", "ESI"):
        headers = ["Episode ID", "Employee Number", "Employee Name",
                   "UAN" if record.scheme == "PF" else "ESI Number",
                   "Wages Used", "Employee Contribution", "Employer Contribution"]
        sheet.append(headers)
        run_query = db.query(PayrollRun).filter(PayrollRun.year == record.year, PayrollRun.month == record.month)
        if record.cost_center_id is not None:
            run_query = run_query.filter(PayrollRun.cost_center_id == record.cost_center_id)
        run_ids = [r.id for r in run_query.all()]
        employee_code = SCHEME_COMPONENT_CODES[record.scheme]
        employer_code = EMPLOYER_SCHEME_COMPONENT_CODES[record.scheme]

        if run_ids:
            payslips = db.query(Payslip).filter(Payslip.run_id.in_(run_ids)).all()
            for payslip in payslips:
                lines = db.query(PayslipLine).filter(PayslipLine.payslip_id == payslip.id).all()
                by_code = {l.component_code: l.amount for l in lines}
                employee_amount = by_code.get(employee_code, 0.0)
                if not employee_amount:
                    continue
                employer_amount = by_code.get(employer_code, 0.0)
                episode = payslip.episode
                statutory_info = (
                    db.query(StatutoryInfo)
                    .filter(StatutoryInfo.episode_id == payslip.episode_id, StatutoryInfo.effective_to.is_(None))
                    .order_by(StatutoryInfo.effective_from.desc())
                    .first()
                )
                identifier = (statutory_info.uan if record.scheme == "PF" else statutory_info.esi_number) if statutory_info else None
                sheet.append([
                    episode.id, episode.employee_number,
                    f"{episode.employee.first_name} {episode.employee.last_name}" if episode.employee else None,
                    identifier, payslip.gross_earnings, employee_amount, employer_amount,
                ])
    elif record.scheme == "PT":
        headers = ["Employee Number", "Employee Name", "Gross Earnings", "PT Amount"]
        sheet.append(headers)
        run_query = db.query(PayrollRun).filter(PayrollRun.year == record.year, PayrollRun.month == record.month)
        if record.cost_center_id is not None:
            run_query = run_query.filter(PayrollRun.cost_center_id == record.cost_center_id)
        run_ids = [r.id for r in run_query.all()]
        if run_ids:
            payslips = db.query(Payslip).filter(Payslip.run_id.in_(run_ids)).all()
            for payslip in payslips:
                line = (
                    db.query(PayslipLine)
                    .filter(PayslipLine.payslip_id == payslip.id, PayslipLine.component_code == "PT", PayslipLine.amount != 0)
                    .first()
                )
                if not line:
                    continue
                episode = payslip.episode
                sheet.append([
                    episode.employee_number,
                    f"{episode.employee.first_name} {episode.employee.last_name}" if episode.employee else None,
                    payslip.gross_earnings, line.amount,
                ])
    elif record.scheme == "LWF":
        headers = ["Employee Number", "Employee Name", "Employee Amount", "Employer Amount"]
        sheet.append(headers)
        run_query = db.query(PayrollRun).filter(PayrollRun.year == record.year, PayrollRun.month == record.month)
        if record.cost_center_id is not None:
            run_query = run_query.filter(PayrollRun.cost_center_id == record.cost_center_id)
        run_ids = [r.id for r in run_query.all()]
        if run_ids:
            payslips = db.query(Payslip).filter(Payslip.run_id.in_(run_ids)).all()
            for payslip in payslips:
                lines = db.query(PayslipLine).filter(PayslipLine.payslip_id == payslip.id).all()
                by_code = {l.component_code: l.amount for l in lines}
                employee_amount = by_code.get("LWF", 0.0)
                if not employee_amount:
                    continue
                episode = payslip.episode
                sheet.append([
                    episode.employee_number,
                    f"{episode.employee.first_name} {episode.employee.last_name}" if episode.employee else None,
                    employee_amount, by_code.get("EMPLOYER_LWF", 0.0),
                ])
    elif record.scheme == "GRATUITY":
        headers = ["Employee Number", "Employee Name", "Years of Service", "Gratuity Amount"]
        sheet.append(headers)
        _, last_day = calendar.monthrange(record.year, record.month)
        month_start = dt.datetime(record.year, record.month, 1)
        month_end = dt.datetime(record.year, record.month, last_day, 23, 59, 59)
        settlements = (
            db.query(FullFinalSettlement)
            .filter(
                FullFinalSettlement.status == "COMPLETED",
                FullFinalSettlement.processed_at >= month_start,
                FullFinalSettlement.processed_at <= month_end,
                FullFinalSettlement.gratuity_amount > 0,
            )
            .all()
        )
        if record.cost_center_id is not None:
            settlements = [
                s for s in settlements
                if approval_service.current_cost_center_id(db, s.episode_id) == record.cost_center_id
            ]
        for settlement in settlements:
            episode = db.query(EmploymentEpisode).filter(EmploymentEpisode.id == settlement.episode_id).first()
            sheet.append([
                episode.employee_number if episode else None,
                f"{episode.employee.first_name} {episode.employee.last_name}" if episode and episode.employee else None,
                settlement.gratuity_years_of_service, settlement.gratuity_amount,
            ])

    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for col_idx in range(1, sheet.max_column + 1):
        sheet.column_dimensions[sheet.cell(row=1, column=col_idx).column_letter].width = 22

    return wb
