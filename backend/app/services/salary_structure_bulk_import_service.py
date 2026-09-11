"""Bulk salary-structure assignment + ad-hoc pay entry upload via .xlsx -
mirrors bulk_import_service.py/attendance_bulk_import_service.py's exact
shape (COLUMNS/SAMPLE_ROW, template+reference-sheet download, per-row
db.begin_nested() savepoints)."""
import io
from datetime import date, datetime

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session

from app.models.models import AdhocPayEntry, EmploymentEpisode, SalaryComponent, User
from app.services import payroll_service

# ---------------------------------------------------------------------------
# Salary structure bulk assignment
# ---------------------------------------------------------------------------

STRUCTURE_COLUMNS = [
    ("employee_number", "Employee Number*"),
    ("component_code", "Component Code*"),
    ("amount", "Amount"),
    ("percentage", "Percentage"),
    ("effective_from", "Effective From* (YYYY-MM-DD)"),
]

STRUCTURE_SAMPLE_ROW = {
    "employee_number": "EMP00200", "component_code": "BASIC", "amount": "20000", "percentage": "",
    "effective_from": "2026-01-01",
}


def build_structure_template_workbook(db: Session) -> Workbook:
    wb = Workbook()
    sheet = wb.active
    sheet.title = "Salary Structure"

    headers = [header for _, header in STRUCTURE_COLUMNS]
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    sheet.append([STRUCTURE_SAMPLE_ROW.get(key, "") for key, _ in STRUCTURE_COLUMNS])

    for col_idx, _ in enumerate(STRUCTURE_COLUMNS, start=1):
        sheet.column_dimensions[sheet.cell(row=1, column=col_idx).column_letter].width = 26

    ref = wb.create_sheet("Reference Values")
    ref["A1"] = "Component Codes"
    ref["A1"].font = Font(bold=True)
    for i, component in enumerate(db.query(SalaryComponent).filter(SalaryComponent.is_active.is_(True)).all(), start=2):
        ref[f"A{i}"] = component.code
    ref.column_dimensions["A"].width = 26

    return wb


def _cell_str(value) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def _cell_date(value) -> date | None:
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    try:
        return datetime.strptime(str(value).strip(), "%Y-%m-%d").date()
    except ValueError:
        return None


def _cell_float(value) -> float | None:
    if value is None or value == "":
        return None
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


def import_structure_workbook(db: Session, file_bytes: bytes, actor: User) -> dict:
    wb = load_workbook(io.BytesIO(file_bytes), data_only=True)
    sheet = wb["Salary Structure"] if "Salary Structure" in wb.sheetnames else wb.active

    header_row = [(_cell_str(c.value) or "") for c in sheet[1]]
    key_by_column_index = {}
    for idx, header in enumerate(header_row):
        for key, expected_header in STRUCTURE_COLUMNS:
            if header.strip().lower() == expected_header.strip().lower():
                key_by_column_index[idx] = key
                break

    created = 0
    errors = []

    for row_number, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
        if row is None or all(v is None or str(v).strip() == "" for v in row):
            continue

        data = {}
        for idx, value in enumerate(row):
            key = key_by_column_index.get(idx)
            if key:
                data[key] = value

        employee_number = _cell_str(data.get("employee_number"))
        component_code = _cell_str(data.get("component_code"))
        effective_from = _cell_date(data.get("effective_from"))

        if not employee_number or not component_code or not effective_from:
            errors.append({"row": row_number, "message": "Employee Number, Component Code and Effective From are required"})
            continue

        savepoint = db.begin_nested()
        try:
            episode = db.query(EmploymentEpisode).filter(EmploymentEpisode.employee_number == employee_number).first()
            if not episode:
                raise ValueError(f"Employee Number '{employee_number}' not found")

            component = db.query(SalaryComponent).filter(SalaryComponent.code == component_code).first()
            if not component:
                raise ValueError(f"Component Code '{component_code}' not found")

            amount = _cell_float(data.get("amount"))
            percentage = _cell_float(data.get("percentage"))
            if amount is None and percentage is None:
                raise ValueError("Either Amount or Percentage must be provided")

            payroll_service.set_salary_structure_component(db, episode.id, component.id, amount, percentage, effective_from, actor)

            savepoint.commit()
            created += 1
        except (ValueError, IntegrityError) as exc:
            savepoint.rollback()
            errors.append({"row": row_number, "message": str(exc.__cause__ or exc) if isinstance(exc, IntegrityError) else str(exc)})

    return {"created": created, "errors": errors}


# ---------------------------------------------------------------------------
# Ad-hoc pay entry bulk upload
# ---------------------------------------------------------------------------

ADHOC_COLUMNS = [
    ("employee_number", "Employee Number*"),
    ("year", "Year*"),
    ("month", "Month* (1-12)"),
    ("label", "Label*"),
    ("amount", "Amount*"),
    ("is_earning", "Is Earning (YES/NO)"),
    ("remarks", "Remarks"),
]

ADHOC_SAMPLE_ROW = {
    "employee_number": "EMP00200", "year": "2026", "month": "1", "label": "Festival Bonus",
    "amount": "1000", "is_earning": "YES", "remarks": "",
}


def build_adhoc_template_workbook() -> Workbook:
    wb = Workbook()
    sheet = wb.active
    sheet.title = "Adhoc Entries"

    headers = [header for _, header in ADHOC_COLUMNS]
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    sheet.append([ADHOC_SAMPLE_ROW.get(key, "") for key, _ in ADHOC_COLUMNS])

    for col_idx, _ in enumerate(ADHOC_COLUMNS, start=1):
        sheet.column_dimensions[sheet.cell(row=1, column=col_idx).column_letter].width = 26

    return wb


def import_adhoc_workbook(db: Session, file_bytes: bytes, actor: User) -> dict:
    wb = load_workbook(io.BytesIO(file_bytes), data_only=True)
    sheet = wb["Adhoc Entries"] if "Adhoc Entries" in wb.sheetnames else wb.active

    header_row = [(_cell_str(c.value) or "") for c in sheet[1]]
    key_by_column_index = {}
    for idx, header in enumerate(header_row):
        for key, expected_header in ADHOC_COLUMNS:
            if header.strip().lower() == expected_header.strip().lower():
                key_by_column_index[idx] = key
                break

    created = 0
    errors = []

    for row_number, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
        if row is None or all(v is None or str(v).strip() == "" for v in row):
            continue

        data = {}
        for idx, value in enumerate(row):
            key = key_by_column_index.get(idx)
            if key:
                data[key] = value

        employee_number = _cell_str(data.get("employee_number"))
        label = _cell_str(data.get("label"))
        amount = _cell_float(data.get("amount"))

        if not employee_number or not label or amount is None:
            errors.append({"row": row_number, "message": "Employee Number, Label and Amount are required"})
            continue

        savepoint = db.begin_nested()
        try:
            episode = db.query(EmploymentEpisode).filter(EmploymentEpisode.employee_number == employee_number).first()
            if not episode:
                raise ValueError(f"Employee Number '{employee_number}' not found")

            year = int(_cell_float(data.get("year")) or 0)
            month = int(_cell_float(data.get("month")) or 0)
            if not year or not (1 <= month <= 12):
                raise ValueError("Year and Month (1-12) are required")

            is_earning_raw = (_cell_str(data.get("is_earning")) or "YES").strip().upper()
            is_earning = is_earning_raw != "NO"

            entry = AdhocPayEntry(
                episode_id=episode.id, year=year, month=month, label=label, amount=amount,
                is_earning=is_earning, remarks=_cell_str(data.get("remarks")), created_by_id=actor.id,
            )
            db.add(entry)

            savepoint.commit()
            created += 1
        except (ValueError, IntegrityError) as exc:
            savepoint.rollback()
            errors.append({"row": row_number, "message": str(exc.__cause__ or exc) if isinstance(exc, IntegrityError) else str(exc)})

    return {"created": created, "errors": errors}
