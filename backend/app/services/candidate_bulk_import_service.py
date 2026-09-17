"""Bulk candidate upload via .xlsx (Module 5 recruitment - batch entry
point mirroring bulk_import_service.py's Employee upload). Unlike
Employee bulk upload, this is create-only: a Candidate has no natural
external key that exists before insert (reference_number is generated
from the row's own id after the fact, the way a human "New Candidate"
submission works), so there is no update-by-match path here - every
valid row always creates a new Candidate. Documents/Driving Licence
requirement resolution stay per-candidate, filled in later on the
Candidate Detail page, same reasoning as Employee bulk upload skipping
file-upload fields.
"""
import io
from datetime import date, datetime

from fastapi import HTTPException
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session

from app.core.validators import validate_aadhaar, validate_mobile, validate_pan
from app.models.enums import AuditAction
from app.models.models import Candidate, CostCenter, Designation, EmployeeCategory, Project, User
from app.services import audit_service, recruitment_service

COLUMNS = [
    ("first_name", "First Name*"),
    ("middle_name", "Middle Name"),
    ("last_name", "Last Name*"),
    ("father_husband_name", "Father's/Husband's Name"),
    ("gender", "Gender (MALE/FEMALE/OTHER)"),
    ("date_of_birth", "Date of Birth (YYYY-MM-DD)"),
    ("mobile_number", "Mobile Number"),
    ("alternate_mobile_number", "Alternate Mobile Number"),
    ("personal_email", "Personal Email"),
    ("educational_qualification", "Educational Qualification"),
    ("current_designation", "Current Designation"),
    ("current_company_name", "Current Company Name"),
    ("current_company_details", "Current Company Details"),
    ("current_date_of_joining", "Date of Joining (Current Company) (YYYY-MM-DD)"),
    ("total_experience_years", "Total Experience (years)"),
    ("aadhaar", "Aadhaar"),
    ("aadhaar_name", "Name (as on Aadhaar)"),
    ("aadhaar_dob", "Date of Birth (as on Aadhaar) (YYYY-MM-DD)"),
    ("pan", "PAN"),
    ("pan_name", "Name (as on PAN)"),
    ("pan_dob", "Date of Birth (as on PAN) (YYYY-MM-DD)"),
    ("dl_licence_number", "Driving Licence Number"),
    ("dl_badge_number", "Driving Licence Badge Number"),
    ("dl_vehicle_class", "Driving Licence Vehicle Class"),
    ("dl_issuing_authority", "Driving Licence Issuing Authority"),
    ("dl_issue_date", "Driving Licence Issue Date (YYYY-MM-DD)"),
    ("dl_expiry_date", "Driving Licence Expiry Date (YYYY-MM-DD)"),
    ("applied_designation", "Applied Designation* (must match Organization Setup)"),
    ("applied_employee_category", "Applied Employee Category (must match Organization Setup)"),
    ("applied_cost_center", "Applied Cost Center* (must match Organization Setup)"),
    ("applied_project", "Applied Project (must match Organization Setup)"),
    ("applied_date", "Applied Date (YYYY-MM-DD)"),
    ("source", "Source"),
    ("remarks", "Remarks"),
]

SAMPLE_ROW = {
    "first_name": "Anitha", "middle_name": "", "last_name": "Raman", "father_husband_name": "Ramaswamy",
    "gender": "FEMALE", "date_of_birth": "1996-08-21",
    "mobile_number": "9876512345", "alternate_mobile_number": "", "personal_email": "anitha.r@example.com",
    "educational_qualification": "Diploma",
    "current_designation": "Driver", "current_company_name": "XYZ Travels",
    "current_company_details": "Regional bus operator", "current_date_of_joining": "2020-04-01",
    "total_experience_years": 4.2,
    "aadhaar": "234123412342", "aadhaar_name": "Anitha Raman", "aadhaar_dob": "1996-08-21",
    "pan": "ABCDE1234G", "pan_name": "Anitha Raman", "pan_dob": "1996-08-21",
    "dl_licence_number": "", "dl_badge_number": "", "dl_vehicle_class": "",
    "dl_issuing_authority": "", "dl_issue_date": "", "dl_expiry_date": "",
    "applied_designation": "Bus Driver", "applied_employee_category": "Driver",
    "applied_cost_center": "Puducherry", "applied_project": "",
    "applied_date": "2026-01-01", "source": "Referral", "remarks": "",
}


def build_template_workbook(db: Session) -> Workbook:
    wb = Workbook()
    sheet = wb.active
    sheet.title = "Candidates"

    headers = [header for _, header in COLUMNS]
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = Font(bold=True)

    sheet.append([SAMPLE_ROW.get(key, "") for key, _ in COLUMNS])
    for col_idx, _ in enumerate(COLUMNS, start=1):
        sheet.column_dimensions[sheet.cell(row=1, column=col_idx).column_letter].width = 26

    ref = wb.create_sheet("Reference Values")
    ref["A1"] = "Use these exact names in the matching columns on the Candidates sheet"
    ref["A1"].font = Font(bold=True)
    ref.delete_rows(1)

    def _dump(row_label, col_letter, names):
        ref[f"{col_letter}1"] = row_label
        ref[f"{col_letter}1"].font = Font(bold=True)
        for i, name in enumerate(names, start=2):
            ref[f"{col_letter}{i}"] = name
        ref.column_dimensions[col_letter].width = 26

    _dump("Designations", "A", [d.name for d in db.query(Designation).filter(Designation.is_active.is_(True)).all()])
    _dump("Employee Categories", "B", [c.name for c in db.query(EmployeeCategory).filter(EmployeeCategory.is_active.is_(True)).all()])
    _dump("Cost Centers", "C", [c.name for c in db.query(CostCenter).filter(CostCenter.is_active.is_(True)).all()])
    _dump("Projects", "D", [p.name for p in db.query(Project).filter(Project.is_active.is_(True)).all()])

    return wb


def _cell_str(value) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def _cell_str_upper(value) -> str | None:
    """Same "everything in capital letters" convention as
    bulk_import_service._cell_str_upper - not used for values passed to
    _lookup() (master data names matched case-sensitively) or for
    personal_email (case can be meaningful)."""
    text = _cell_str(value)
    return text.upper() if text else None


def _cell_date(value) -> date | None:
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    try:
        return datetime.strptime(str(value).strip(), "%Y-%m-%d").date()
    except ValueError:
        return None


def _cell_float(value) -> float | None:
    if value is None or value == "":
        return None
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


def _lookup(db: Session, model, name: str | None):
    if not name:
        return None
    return db.query(model).filter(model.name == name, model.is_active.is_(True)).first()


def import_candidates_workbook(db: Session, file_bytes: bytes, actor: User) -> dict:
    wb = load_workbook(io.BytesIO(file_bytes), data_only=True)
    sheet = wb["Candidates"] if "Candidates" in wb.sheetnames else wb.active

    header_row = [(_cell_str(c.value) or "") for c in sheet[1]]
    key_by_column_index = {}
    for idx, header in enumerate(header_row):
        for key, expected_header in COLUMNS:
            if header.strip().lower() == expected_header.strip().lower():
                key_by_column_index[idx] = key
                break

    created = 0
    errors = []

    for row_number, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
        if row is None or all(v is None or str(v).strip() == "" for v in row):
            continue  # skip blank rows

        data = {}
        for idx, value in enumerate(row):
            key = key_by_column_index.get(idx)
            if key:
                data[key] = value

        first_name = _cell_str_upper(data.get("first_name"))
        last_name = _cell_str_upper(data.get("last_name"))
        applied_designation_name = _cell_str(data.get("applied_designation"))
        applied_cost_center_name = _cell_str(data.get("applied_cost_center"))

        if not first_name or not last_name:
            errors.append({"row": row_number, "message": "First Name and Last Name are required"})
            continue

        designation = _lookup(db, Designation, applied_designation_name)
        if not designation:
            errors.append({"row": row_number, "message": f"Applied Designation '{applied_designation_name or ''}' not found in Organization Setup"})
            continue

        cost_center = _lookup(db, CostCenter, applied_cost_center_name)
        if not cost_center:
            errors.append({"row": row_number, "message": f"Applied Cost Center '{applied_cost_center_name or ''}' not found in Organization Setup"})
            continue

        employee_category = _lookup(db, EmployeeCategory, _cell_str(data.get("applied_employee_category")))
        project = _lookup(db, Project, _cell_str(data.get("applied_project")))

        savepoint = db.begin_nested()
        try:
            aadhaar = _cell_str(data.get("aadhaar"))
            if aadhaar:
                aadhaar = validate_aadhaar(aadhaar)
            pan = _cell_str(data.get("pan"))
            if pan:
                pan = validate_pan(pan)
            mobile_number = _cell_str(data.get("mobile_number"))
            if mobile_number:
                mobile_number = validate_mobile(mobile_number)
            alternate_mobile_number = _cell_str(data.get("alternate_mobile_number"))
            if alternate_mobile_number:
                alternate_mobile_number = validate_mobile(alternate_mobile_number)
            dl_licence_number = _cell_str(data.get("dl_licence_number"))

            try:
                recruitment_service.validate_unique_identifiers(db, aadhaar, pan, dl_licence_number)
            except HTTPException as exc:
                raise ValueError(exc.detail) from exc

            candidate = Candidate(
                first_name=first_name, middle_name=_cell_str_upper(data.get("middle_name")), last_name=last_name,
                father_husband_name=_cell_str_upper(data.get("father_husband_name")),
                gender=_cell_str_upper(data.get("gender")), date_of_birth=_cell_date(data.get("date_of_birth")),
                mobile_number=mobile_number, alternate_mobile_number=alternate_mobile_number,
                personal_email=_cell_str(data.get("personal_email")),
                educational_qualification=_cell_str_upper(data.get("educational_qualification")),
                current_designation=_cell_str_upper(data.get("current_designation")),
                current_company_name=_cell_str_upper(data.get("current_company_name")),
                current_company_details=_cell_str_upper(data.get("current_company_details")),
                current_date_of_joining=_cell_date(data.get("current_date_of_joining")),
                total_experience_years=_cell_float(data.get("total_experience_years")),
                aadhaar=aadhaar, aadhaar_name=_cell_str_upper(data.get("aadhaar_name")), aadhaar_dob=_cell_date(data.get("aadhaar_dob")),
                pan=pan, pan_name=_cell_str_upper(data.get("pan_name")), pan_dob=_cell_date(data.get("pan_dob")),
                dl_licence_number=dl_licence_number, dl_badge_number=_cell_str_upper(data.get("dl_badge_number")),
                dl_vehicle_class=_cell_str_upper(data.get("dl_vehicle_class")),
                dl_issuing_authority=_cell_str_upper(data.get("dl_issuing_authority")),
                dl_issue_date=_cell_date(data.get("dl_issue_date")), dl_expiry_date=_cell_date(data.get("dl_expiry_date")),
                applied_designation_id=designation.id,
                applied_employee_category_id=employee_category.id if employee_category else None,
                applied_cost_center_id=cost_center.id,
                applied_project_id=project.id if project else None,
                applied_date=_cell_date(data.get("applied_date")) or date.today(),
                source=_cell_str_upper(data.get("source")), remarks=_cell_str(data.get("remarks")),
            )
            db.add(candidate)
            db.flush()
            candidate.reference_number = recruitment_service.generate_reference_number(db, candidate)
            db.add(candidate)
            audit_service.record(db, "CANDIDATE", candidate.id, AuditAction.CREATE, actor, new_value=f"{candidate.reference_number} (bulk upload)")
            savepoint.commit()
            created += 1
        except (ValueError, IntegrityError) as exc:
            savepoint.rollback()
            errors.append({"row": row_number, "message": str(exc.__cause__ or exc) if isinstance(exc, IntegrityError) else str(exc)})

    return {"created": created, "errors": errors}
